"""Extracción de tablas desde CSV (stdlib), Excel (openpyxl para .xlsx y xlrd
para .xls) y PDF (CF2-03, CF2-04, soporte de cartolas Itaú)."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time

from cuentafaro.importing.contract import ImportFormatError


@dataclass
class ExtractedTable:
    headers: list[str]
    rows: list[dict[str, str]]
    first_row_number: int
    separator: str | None = None
    default_year: int | None = None
    preamble: list[dict[str, str]] = field(default_factory=list)


SUPPORTED_KINDS = ("csv", "excel", "pdf")


def guess_source_kind(filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return "csv"
    if lower.endswith((".xlsx", ".xlsm", ".xls", ".xltx")):
        return "excel"
    if lower.endswith(".pdf"):
        return "pdf"
    raise ImportFormatError(f"formato no soportado; use .csv, .xlsx o .pdf: {filename}")


def extract_table(
    content: bytes,
    *,
    source_kind: str,
    header_row: int = 1,
    separator: str | None = None,
) -> ExtractedTable:
    if header_row < 1:
        raise ImportFormatError("header_row debe ser mayor o igual a 1")
    if source_kind == "csv":
        return _extract_csv(content, header_row=header_row, separator=separator)
    if source_kind == "excel":
        return _extract_excel(content, header_row=header_row)
    if source_kind == "pdf":
        return _extract_pdf(content)
    raise ImportFormatError(f"formato desconocido: {source_kind}")


def _extract_pdf(content: bytes) -> ExtractedTable:
    from cuentafaro.ai.ocr import extract_pdf_text_as_rows, pdf_rows_to_table

    rows = extract_pdf_text_as_rows(content)
    if not rows:
        raise ImportFormatError("no se extrajeron filas del PDF")
    return pdf_rows_to_table(rows)


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1", errors="replace")


def _extract_csv(content: bytes, *, header_row: int, separator: str | None) -> ExtractedTable:
    text = _decode(content)
    kwargs = {"delimiter": separator} if separator else _sniff_delimiter(text)
    rows = list(csv.reader(io.StringIO(text), **kwargs))
    if header_row - 1 >= len(rows):
        raise ImportFormatError("la fila de encabezados no existe en el archivo")
    headers = _clean_headers(rows[header_row - 1])
    data = rows[header_row:]
    table = _build_table(
        headers,
        data,
        first_row_number=header_row + 1,
        separator=kwargs.get("delimiter"),
    )
    table.preamble = _preamble_rows(rows[: header_row - 1])
    return table


def _sniff_delimiter(text: str) -> dict:
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=";,\t")
        return {"delimiter": dialect.delimiter}
    except csv.Error:
        return {"delimiter": ","}


def _extract_excel(content: bytes, *, header_row: int) -> ExtractedTable:
    records = _excel_records(content)
    if not records:
        raise ImportFormatError("no se encontraron filas en el Excel")
    header_row = _locate_header_row(records, header_row)
    if header_row - 1 >= len(records):
        raise ImportFormatError("la fila de encabezados no existe en el Excel")
    headers = _clean_headers(records[header_row - 1])
    data = [(_cell_to_text(cell) for cell in row) for row in records[header_row:]]
    data = [tuple(cell for cell in cells) for cells in data]
    table = _build_table(headers, data, first_row_number=header_row + 1)
    table.default_year = _detect_year(records[: min(header_row + 20, len(records))])
    table.preamble = _preamble_rows(records[: header_row - 1])
    return table


def _preamble_rows(records: list) -> list[dict[str, str]]:
    """Convierte las filas previas a la tabla en dicts ``columna_N`` (CF-tarjetas).

    Las cartolas de tarjeta ponen el resumen (Total a pagar, Pago mínimo,
    vencimiento, cupo) sobre la tabla de movimientos; se conservan para que el
    detector de estado de cuenta pueda leerlos sin tocar el parser.
    """
    rows: list[dict[str, str]] = []
    for record in records:
        values = [str(_cell_to_text(cell)).strip() for cell in record]
        row = {
            f"columna_{index + 1}": value for index, value in enumerate(values) if value
        }
        if row:
            rows.append(row)
    return rows


_XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_YEAR_PATTERN = re.compile(r"\b(20\d{2})\b")


def _excel_records(content: bytes) -> list[tuple]:
    """Lee filas de valores desde un Excel viejo (.xls, OLE2) o moderno (.xlsx)."""
    if content[:8] == _XLS_MAGIC:
        from xlrd import XL_CELL_BOOLEAN, XL_CELL_DATE, XL_CELL_NUMBER, open_workbook

        try:
            book = open_workbook(file_contents=content)
        except Exception as error:
            raise ImportFormatError(f"no se pudo leer el Excel: {error}") from error
        sheet = book.sheet_by_index(0)
        records: list[tuple] = []
        for row_index in range(sheet.nrows):
            row: list[object] = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                if cell.ctype == XL_CELL_DATE:
                    row.append(datetime_from_xls(cell.value, book.datemode))
                elif cell.ctype == XL_CELL_NUMBER:
                    row.append(cell.value)
                elif cell.ctype == XL_CELL_BOOLEAN:
                    row.append(bool(cell.value))
                else:
                    row.append(cell.value or None)
            records.append(tuple(row))
        return records
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:  # archivo corrupto o no es Excel
        raise ImportFormatError(f"no se pudo leer el Excel: {error}") from error
    sheet = workbook.active
    records = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return records


def datetime_from_xls(value: float, datemode: int) -> object:
    """Convierte un número serial xlrd a una fecha (CF4/Itaú: cartolas .xls)."""
    try:
        import xlrd

        return xlrd.xldate.xldate_as_datetime(value, datemode).date()
    except (ValueError, TypeError, OverflowError):
        return value


_HEADER_ALIASES = (
    "fecha",
    "fechamovimiento",
    "fechaoperacion",
    "descripcion",
    "detalle",
    "glosa",
    "concepto",
    "comercio",
    "movimiento",
    "cargo",
    "abono",
    "debito",
    "credito",
    "saldo",
    "operacion",
    "referencia",
    "folio",
    "monto",
    "importe",
    "valor",
)


def _normalize_header(value: object) -> str:
    import unicodedata

    text = _cell_to_text(value).strip().lower()
    text = "".join(c for c in unicodedata.normalize("NFD", text) if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", text)


def _alias_hits(row: tuple) -> int:
    hits = 0
    for cell in row[:12]:
        name = _normalize_header(cell)
        if name and any(name.endswith(alias) or alias in name for alias in _HEADER_ALIASES):
            hits += 1
    return hits


def _locate_header_row(records: list[tuple], header_row: int) -> int:
    """Devuelve la fila (1-based) de encabezados, buscando la tabla si la fila
    indicada cae en un bloque de metadatos (cartolas con cabecera). Elige la
    fila candidata con más coincidencias de alias, desempatando por la más alta."""
    limit = min(max(header_row + 40, 41), len(records))
    best, best_hits = header_row - 1, _alias_hits(records[header_row - 1])
    for index in range(header_row - 1, limit):
        hits = _alias_hits(records[index])
        if hits > best_hits:
            best, best_hits = index, hits
    if best_hits < 2:
        return header_row
    return best + 1


def _detect_year(records: list[tuple]) -> int | None:
    from collections import Counter

    years: Counter[int] = Counter()
    for row in records:
        for cell in row:
            match = _YEAR_PATTERN.search(str(cell))
            if match:
                years[int(match.group(1))] += 1
    if not years:
        return None
    return years.most_common(1)[0][0]


def _clean_headers(header_row: tuple) -> list[str]:
    headers = []
    for index, cell in enumerate(header_row):
        text = _cell_to_text(cell).strip()
        headers.append(text or f"columna_{index + 1}")
    if not any(header for header in headers):
        raise ImportFormatError("los encabezados están vacíos")
    if len(headers) != len(set(headers)):
        raise ImportFormatError("hay encabezados duplicados en el archivo")
    return headers


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _build_table(
    headers: list[str],
    data: list[tuple],
    *,
    first_row_number: int,
    separator: str | None = None,
) -> ExtractedTable:
    rows: list[dict[str, str]] = []
    for cells in data:
        values: dict[str, str] = {}
        empty = True
        for index, name in enumerate(headers):
            cell = cells[index] if index < len(cells) else None
            text = _cell_to_text(cell).strip()
            if text:
                empty = False
            values[name] = text
        if not empty:
            rows.append(values)
    return ExtractedTable(
        headers=headers,
        rows=rows,
        first_row_number=first_row_number,
        separator=separator,
    )
